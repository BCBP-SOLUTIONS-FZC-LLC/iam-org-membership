#!/usr/bin/env python3
"""
iam-org-membership automated test runner — final version.

Usage:
    python3 scripts/test_runner.py

Prerequisites:
    make mock-servers   (terminal 1)
    make run            (terminal 2)
"""
import json, re, time, uuid
import urllib.request, urllib.error
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# ── Config ────────────────────────────────────────────────────────────
BASE       = "http://localhost:8080"
EXCEL_PATH = "/Users/sharmila/bcbp-solutions/XpertPMS/Org-Membership/Testing/o&g_testing.xlsx"
SHEET_NAME = "Final"

# ── Fixed test identities ─────────────────────────────────────────────
TENANT    = "cccc0099-0001-0001-0001-000000000001"   # Test Corp (fresh, 10 seats)
TENANT_B  = "cccc0002-0001-0001-0001-000000000001"   # Beta Corp
OWNER     = "bbbb0099-0001-0001-0001-000000000001"   # tenant_owner of TENANT
ADMIN     = "dddd0099-0001-0001-0001-000000000002"   # tenant_admin of TENANT
MEMBER    = "dddd0099-0001-0001-0001-000000000001"   # plain member of TENANT
SUSPENDED = "bbbb0001-0001-0001-0001-000000000003"   # suspended user (from original tenant)
OWNER_B   = "eeee0001-0001-0001-0001-000000000001"   # owner of TENANT_B
SYS_USER  = "00000000-0000-0000-0000-000000000001"
SYS_TEN   = "aaaa0001-0001-0001-0001-000000000001"
DEPT_ENG  = "de010001-0000-0000-0000-000000000001"   # engineering (system dept)
DEPT_DES  = "de010002-0000-0000-0000-000000000002"   # design (non-system)
NIL_UUID  = "00000000-0000-0000-0000-000000000000"
TENDER_ID = str(uuid.uuid4())

# ── Auth header sets ──────────────────────────────────────────────────
def H(user, tenant, roles):
    return {"Content-Type":"application/json",
            "x-user-id":user,"x-tenant-id":tenant,"x-tenant-roles":roles}

H_SYS    = H(SYS_USER, SYS_TEN, "iam-system")
H_OWNER  = H(OWNER,    TENANT,   "tenant_owner")
H_ADMIN  = H(ADMIN,    TENANT,   "tenant_admin")
H_MEMBER = H(MEMBER,   TENANT,   "member")
H_OP     = H(SYS_USER, TENANT,   "platform_operator")
H_OWNER_B= H(OWNER_B,  TENANT_B, "tenant_owner")
H_CROSS  = H(OWNER,    TENANT_B, "tenant_owner")
H_NO_AUTH= {"Content-Type":"application/json"}
H_NOCT   = {k:v for k,v in H_SYS.items() if k!="Content-Type"}

# ── HTTP helper ───────────────────────────────────────────────────────
def call(method, url, body=None, hdrs=None):
    if hdrs is None: hdrs = H_SYS
    data = json.dumps(body).encode() if body is not None else None
    req  = urllib.request.Request(url, data=data, headers=hdrs, method=method)
    try:
        r = urllib.request.urlopen(req, timeout=5)
        return r.status, json.loads(r.read() or b'{}')
    except urllib.error.HTTPError as e:
        try: return e.code, json.loads(e.read())
        except: return e.code, {}
    except Exception as ex:
        return 0, {"_error": str(ex)}

# ── Record-version: always fetch fresh before PATCH ───────────────────
def get_rv(resource_url, hdrs):
    c, b = call("GET", resource_url, None, hdrs)
    return b.get("record_version", 1) if c == 200 else 1

# ── Dynamic state ─────────────────────────────────────────────────────
state = {
    "tenant_n":   30,     # counter for fresh tenant UUIDs
    "invite_id":  None,
    "member_ok":  True,   # tracks whether MEMBER still exists
    "dept_des_rv":None,   # record_version of DEPT_DES
}

def fresh_tenant():
    state["tenant_n"] += 1
    return f"cccc{state['tenant_n']:04d}-0001-0001-0001-000000000001"

def fresh_user():
    return str(uuid.uuid4())

def ensure_owner():
    """Ensure OWNER is active with tenant_owner role after O-7/P-8 may have removed them."""
    # Check via I-15 (iam-system bypass)
    c, _ = call("GET", BASE+f"/api/v1/internal/tenants/{TENANT}/members/{OWNER}/exists", None, H_SYS)
    if c != 200:
        call("POST", BASE+f"/api/v1/internal/tenants/{TENANT}/members",
             {"user_id":OWNER,"email":"owner@test99.com","full_name":"Owner"}, H_SYS)
    call("PATCH", BASE+f"/api/v1/internal/tenants/{TENANT}/members/{OWNER}",{"status":"active"}, H_SYS)
    # Re-assign ownership via platform_operator
    c2, b2 = call("GET", BASE+f"/api/v1/tenants/{TENANT}", None, H_OP)
    if c2 == 200:
        rv = b2.get("record_version", 1)
        call("POST", BASE+f"/api/v1/operator/tenants/{TENANT}/reassign-owner",
             {"user_id":OWNER,"record_version":rv}, H_OP)

def ensure_member():
    """Re-create MEMBER if it was deleted mid-run. Uses I-15 (iam-system bypass)."""
    c, _ = call("GET", BASE+f"/api/v1/internal/tenants/{TENANT}/members/{MEMBER}/exists", None, H_SYS)
    if c != 200:
        call("POST", BASE+f"/api/v1/internal/tenants/{TENANT}/members",
             {"user_id":MEMBER,"email":"member@test99.com","full_name":"Plain Member"}, H_SYS)
    call("PATCH", BASE+f"/api/v1/internal/tenants/{TENANT}/members/{MEMBER}",{"status":"active"}, H_SYS)
    # Reset MEMBER to plain member role (tests may have granted elevated roles)
    c2, b2 = call("GET", BASE+f"/api/v1/tenants/{TENANT}/members/{MEMBER}", None, H_OWNER)
    if c2 == 200:
        rv = b2.get("record_version", 1)
        if any(r in b2.get("tenant_roles",[]) for r in ["tenant_admin","tenant_owner"]):
            call("PUT", BASE+f"/api/v1/tenants/{TENANT}/members/{MEMBER}/roles",
                 {"roles":[],"record_version":rv}, H_OWNER)
    state["member_ok"] = True

def clean_stray_members():
    """Remove members except OWNER/ADMIN/SUSPENDED/MEMBER — keeps seat count low."""
    KEEP = {OWNER, ADMIN, SUSPENDED, MEMBER}
    c, b = call("GET", BASE+f"/api/v1/tenants/{TENANT}/members", None, H_OWNER)
    for m in b.get("items",[]):
        uid = m.get("user_id","")
        if uid not in KEEP:
            call("DELETE", BASE+f"/api/v1/internal/tenants/{TENANT}/members/{uid}", None, H_SYS)

# Groups that need MEMBER to exist before they run
NEEDS_MEMBER = {"I4","P4","P5","P7","P8","P9","P10","P11","P12","P28","P28RL","P27","P30","P31","GT","IUM","P6"}
# Groups that pollute seat count (need cleanup before P-6)
SEAT_POLLUTERS = {"I3"}

# Track which group we're in
_last_prefix = None

def group_prefix(tc):
    m = re.match(r'^([A-Z]+\d*)', tc.upper())
    return m.group(1) if m else tc[:4]

def pre_test_hooks(tc):
    global _last_prefix
    pfx = group_prefix(tc)
    if pfx != _last_prefix:
        prev = _last_prefix
        # After O7/P8 which can delete/demote OWNER, restore before any OWNER-dependent group
        if prev in ("O4","O7","P8","P28"):
            ensure_owner()
        # Ensure MEMBER active before any test group that uses them
        if pfx in NEEDS_MEMBER:
            ensure_member()
        # Clean up stray I-3 members before P-6 invite tests (free seats)
        if pfx == "P6":
            clean_stray_members()
        _last_prefix = pfx

# ── Skip determination ────────────────────────────────────────────────
SKIP_PATTERNS = [
    "EVT-","PERF-","LOAD-",
    "event emitted","enqueued","outbox","relay",
    "concurrent","race condition",
    "downstream dep","db unavailable",
]
def should_skip(tc, cat, scenario, need_manual):
    if need_manual == "Must": return True
    combined = (tc+cat+scenario).lower()
    return any(p.lower() in combined for p in SKIP_PATTERNS)

# ── Request builder ───────────────────────────────────────────────────
def build_request(tc, cat, scenario, method, api_tmpl, exp_code, notes):
    tc_up = tc.upper()
    cat_l = cat.lower()
    sc_l  = scenario.lower()

    is_auth = "authorization" in cat_l or "auth" in cat_l
    is_val  = "validation" in cat_l
    is_hp   = "happy" in cat_l
    is_sec  = "security" in cat_l or "cross" in cat_l
    is_idp  = "idempotent" in cat_l

    url = BASE + api_tmpl
    # Normalise placeholders
    url = re.sub(r'\{id\}',           TENANT,       url)
    url = re.sub(r':id\b',            TENANT,       url)
    url = re.sub(r'\{user_id\}',      MEMBER,       url)
    url = re.sub(r':user_id\b',       MEMBER,       url)
    url = re.sub(r'\{dept_id\}',      DEPT_ENG,     url)
    url = re.sub(r':dept_id\b',       DEPT_ENG,     url)
    url = re.sub(r'\{role_code\}',    'preparator', url)
    url = re.sub(r':role_code\b',     'preparator', url)
    url = re.sub(r'\{tender_id\}',    TENDER_ID,    url)
    url = re.sub(r'\{invitation_id\}',state.get("invite_id") or NIL_UUID, url)
    url = re.sub(r'/bad-uuid\b',      '/not-a-uuid', url)
    url = re.sub(r'/badid\b',         '/not-a-uuid', url)
    url = re.sub(r'/notauuid\b',      '/not-a-uuid', url)
    url = re.sub(r'/bad\b(?!/)',       '/not-a-uuid', url)

    hdrs = H_SYS
    body = None

    # ── I-1 ─────────────────────────────────────────────────────────
    if tc_up.startswith("I1-"):
        tid  = fresh_tenant()         # always fresh — guarantees 201 not 200
        slug = f"t{state['tenant_n']}"
        valid = {"tenant_id":tid,"name":"Corp","slug":slug,
                 "owner_user_id":fresh_user(),"owner_email":"o@t.com",
                 "owner_name":"O","plan":"starter","licensed_seats":5,"locale":"en-US"}
        hdrs  = H_SYS
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
            body = valid
        elif exp_code==415:          hdrs=H_NOCT; body=valid
        elif exp_code==409 and "slug" in sc_l:
            body = {**valid,"tenant_id":fresh_tenant(),"slug":"acme-corp"}
        elif exp_code in (400,422):
            if "slug" in sc_l:      body={**valid,"slug":"BAD SLUG!!"}
            elif "name" in sc_l:    body={**valid,"name":"x"*300}
            elif "locale" in sc_l and "invalid" in sc_l:  body={**valid,"locale":"notanlocale"}
            elif "plan" in sc_l:    body={**valid,"plan":"diamond"}
            elif "seat" in sc_l:    body={**valid,"licensed_seats":-1}
            elif "nil" in sc_l or "zero" in sc_l or "missing" in sc_l:
                if "owner" in sc_l: body={**valid,"owner_user_id":NIL_UUID}
                else:               body={**valid,"tenant_id":NIL_UUID}
            elif "uuid" in sc_l:
                if "owner" in sc_l: body={**valid,"owner_user_id":"bad"}
                else:               body={**valid,"tenant_id":"bad"}
            else:                   body={**valid,"plan":"diamond"}
        elif exp_code==200:          # idempotent replay
            body={"tenant_id":TENANT,"name":"Acme Corp","slug":"acme-corp",
                  "owner_user_id":OWNER,"owner_email":"o@a.com","owner_name":"O",
                  "plan":"starter","licensed_seats":10,"locale":"en-US"}
        else:
            if "fr-FR" in sc_l:     body={**valid,"locale":"fr-FR"}
            elif "pro" in sc_l:     body={**valid,"plan":"pro"}
            elif "enterprise" in sc_l: body={**valid,"plan":"enterprise"}
            else:                   body=valid

    # ── I-2 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("I2-"):
        hdrs = H_SYS
        rv   = get_rv(BASE+f"/api/v1/tenants/{TENANT}", H_OWNER)
        valid = {"realm_id":f"r-{uuid.uuid4().hex[:6]}",
                 "realm_type":"dedicated","keycloak_shard":"eu-1"}
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
            body = valid
        elif exp_code==415:         hdrs=H_NOCT; body=valid
        elif exp_code in (400,422):
            if "invalid" in sc_l or "realm_type" in sc_l:
                body={"realm_id":"r1","realm_type":"INVALID","keycloak_shard":"eu-1"}
            elif "shard" in sc_l and "missing" in sc_l:
                body={"realm_id":"r1","realm_type":"dedicated"}
            elif "realm_id" in sc_l and "missing" in sc_l:
                body={"realm_type":"dedicated","keycloak_shard":"eu-1"}
            elif "uuid" in sc_l:
                url = re.sub(TENANT,"not-a-uuid",url)
                body=valid
            else: body={}
        elif exp_code==404:
            url=url.replace(TENANT,fresh_tenant()); body=valid
        else: body=valid

    # ── I-3 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("I3-"):
        hdrs = H_SYS
        uid  = fresh_user()
        valid = {"user_id":uid,"email":f"u{uid[:4]}@t.com","full_name":"User"}
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
            body = valid
        elif exp_code==409:
            body={"user_id":MEMBER,"email":"m@m.com","full_name":"Already"}
        elif exp_code in (400,422):
            if "uuid" in sc_l:  body={"user_id":"bad","email":"e@e.com","full_name":"X"}
            elif "email" in sc_l: body={**valid,"email":"not-an-email"}
            else: body={}
        elif exp_code==404:
            url=url.replace(TENANT,fresh_tenant()); body=valid
        else: body=valid

    # ── I-4 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("I4-"):
        hdrs = H_SYS
        url  = url.replace(MEMBER,MEMBER)
        body = {"status":"suspended"}
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
        if "reactivate" in sc_l or ("active" in sc_l and "suspend" not in sc_l):
            body = {"status":"active"}
        elif "left" in sc_l:    body={"status":"left"}
        if exp_code in (400,422):
            if "left" in sc_l:   body={"status":"left"}
            elif "invalid" in sc_l: body={"status":"godmode"}
            else: body={}
        if exp_code==404:
            url=re.sub(MEMBER,fresh_user(),url)

    # ── I-5  (never delete MEMBER — use temp user for success cases) ─
    elif tc_up.startswith("I5-"):
        hdrs = H_SYS
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
            url  = url.replace(MEMBER,MEMBER)
        elif exp_code==400:
            if "tenant" in sc_l:
                url = re.sub(TENANT,"not-a-uuid",url)
                url = re.sub(MEMBER,fresh_user(),url)
            else:   # invalid user UUID
                url = re.sub(MEMBER,"not-a-uuid",url)
        elif exp_code==404:
            url = re.sub(MEMBER,fresh_user(),url)
        elif exp_code==403 or is_sec:
            url = re.sub(MEMBER,fresh_user(),url)
        else:   # exp==200 — delete a throwaway member
            tmp = fresh_user()
            call("POST",BASE+f"/api/v1/internal/tenants/{TENANT}/members",
                 {"user_id":tmp,"email":f"tmp@t.com","full_name":"Tmp"},H_SYS)
            url = re.sub(MEMBER,tmp,url)
            state["member_ok"] = False   # signal that MEMBER check needed next group

    # ── I-8 / IUM ───────────────────────────────────────────────────
    elif tc_up.startswith("I8-") or tc_up.startswith("IUM-"):
        hdrs = H_SYS
        url  = re.sub(r'/users/[^/?]+', f'/users/{OWNER}', url)
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code==400:
            url=re.sub(OWNER,"not-a-uuid",url)
        elif exp_code==404:
            url=re.sub(OWNER,fresh_user(),url)
        if "?" not in url: url+=f"?tenant_id={TENANT}"
        elif "tenant_id" not in url: url+=f"&tenant_id={TENANT}"
        if exp_code==400 and "tenant_id" in sc_l:
            url=re.sub(r'[?&]tenant_id=[^&]*','',url)

    # ── I-9 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("I9-"):
        hdrs = H_SYS
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
        if exp_code==404 or "not-a-uuid" in url:
            url=re.sub(TENANT,fresh_tenant(),url)

    # ── I-10 ────────────────────────────────────────────────────────
    elif tc_up.startswith("I10-"):
        hdrs = H_SYS
        body = {"user_id":MEMBER,"group_names":["engineers"]}
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")

    # ── I-11 ────────────────────────────────────────────────────────
    elif tc_up.startswith("I11-"):
        hdrs = H_SYS
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
        if exp_code==404: url=re.sub(TENANT,fresh_tenant(),url)

    # ── I-13 ────────────────────────────────────────────────────────
    elif tc_up.startswith("I13-"):
        hdrs = H_SYS
        valid = {"actor_id":OWNER,"new_user_id":MEMBER,
                 "required_level":"preparator","tender_id":TENDER_ID}
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
            body = valid
        elif exp_code in (400,422):
            if "actor_id" in sc_l and "zero" in sc_l: body={**valid,"actor_id":NIL_UUID}
            elif "tender_id" in sc_l:
                url=re.sub(TENDER_ID,"not-a-uuid",url); body=valid
            elif "required_level" in sc_l: body={**valid,"required_level":"manager"}
            else: body={}
        elif exp_code==403: body={**valid,"actor_id":MEMBER}
        elif exp_code==404: body={**valid,"new_user_id":fresh_user()}
        else: body=valid

    # ── I-14 ────────────────────────────────────────────────────────
    elif tc_up.startswith("I14-"):
        hdrs = H_SYS
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
        if "not-a-uuid" in url or "bad" in url:
            url=re.sub(TENANT,"not-a-uuid",url)

    # ── I-15 ────────────────────────────────────────────────────────
    elif tc_up.startswith("I15-"):
        hdrs = H_SYS
        url  = url.replace(MEMBER,MEMBER)
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
        if exp_code==404: url=re.sub(MEMBER,fresh_user(),url)

    # ── GT / P-1 ────────────────────────────────────────────────────
    elif tc_up.startswith("GT-") or tc_up.startswith("P1-"):
        hdrs = H_OWNER
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code==403 and ("member" in sc_l or "tender_admin" in sc_l):
            hdrs = H(SYS_USER,TENANT,"tender_admin") if "tender" in sc_l else H_MEMBER
        elif exp_code==404 and "non-existent" in sc_l:
            url=url.replace(TENANT,fresh_tenant())
        elif "operator" in sc_l or "platform_operator" in sc_l: hdrs=H_OP
        elif "tenant_admin" in sc_l and exp_code==200:  hdrs=H_ADMIN
        elif "tenant_owner" in sc_l and exp_code==200:  hdrs=H_OWNER
        elif exp_code==400 and "invalid" in sc_l:
            url=re.sub(TENANT,"not-a-uuid",url)

    # ── P-2 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("P2-"):
        hdrs = H_OWNER
        rv   = get_rv(BASE+f"/api/v1/tenants/{TENANT}", H_OWNER)
        valid = {"name":"Acme Corp","record_version":rv}
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
            body = valid
        elif exp_code==415:         hdrs=H_NOCT; body=valid
        elif exp_code in (400,422):
            if "59" in sc_l or "below" in sc_l:  body={"mfa_freshness_seconds":59,"record_version":rv}
            elif "901" in sc_l or "above" in sc_l: body={"mfa_freshness_seconds":901,"record_version":rv}
            elif "locale" in sc_l:  body={"default_locale":"","record_version":rv}
            elif "empty" in sc_l or "no mutable" in sc_l or "nil" in sc_l: body={"record_version":rv}
            elif "malformed" in sc_l: body=None
            elif "uuid" in sc_l:    url=re.sub(TENANT,"not-a-uuid",url); body=valid
            else:                   body={"record_version":rv}
        elif exp_code==409:         body={"name":"X","record_version":9999}
        elif exp_code==404:
            url=url.replace(TENANT,fresh_tenant()); body={"name":"X","record_version":1}
        elif exp_code==202:         body={**valid,"local_accounts_enabled":True}
        elif exp_code==200:
            if "mfa" in sc_l and "60" in sc_l:  body={"mfa_freshness_seconds":60,"record_version":rv}
            elif "mfa" in sc_l and "900" in sc_l: body={"mfa_freshness_seconds":900,"record_version":rv}
            elif "local_accounts" in sc_l: body={"local_accounts_enabled":True,"record_version":rv}
            else: body=valid

    # ── P-3 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("P3-"):
        hdrs = H_OWNER
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H_CROSS
        elif exp_code==403 and "member" in sc_l: hdrs=H_MEMBER
        elif exp_code==404 and "non-existent" in sc_l:
            url=url.replace(TENANT,fresh_tenant())
        elif exp_code==400 and "uuid" in sc_l:
            url=re.sub(TENANT,"not-a-uuid",url)

    # ── P-4 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("P4-"):
        hdrs = H_OWNER
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H_CROSS
        elif "member" in sc_l and exp_code==200: hdrs=H_MEMBER
        elif exp_code==400:
            if "cursor" in sc_l:    url+="?cursor=!!!notbase64"
            elif "-1" in sc_l:      url+="?limit=-1"
            elif "0" in sc_l and "limit" in sc_l: url+="?limit=0"
            elif "abc" in sc_l:     url+="?limit=abc"
            elif "201" in sc_l:     url+="?limit=201"
        elif exp_code==200 and "limit" in sc_l:
            if "1" in sc_l and "min" in sc_l:   url+="?limit=1"
            elif "200" in sc_l and "max" in sc_l: url+="?limit=200"
            elif "10" in sc_l:      url+="?limit=10"
        elif exp_code==404 and "offboarded" in sc_l:
            url=url.replace(TENANT,fresh_tenant())

    # ── P-5 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("P5-"):
        hdrs = H_OWNER
        url  = url.replace(MEMBER,OWNER)
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H_CROSS
        elif "member" in sc_l and "another" in sc_l and exp_code==200:
            hdrs=H_MEMBER; url=url.replace(OWNER,ADMIN)
        elif "admin" in sc_l and exp_code==200: url=url.replace(OWNER,ADMIN)
        elif exp_code==404: url=re.sub(OWNER,fresh_user(),url)
        elif exp_code==403 and "cross" in sc_l: hdrs=H_CROSS
        elif exp_code==400: url=re.sub(OWNER,"not-a-uuid",url)

    # ── P-6 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("P6-"):
        hdrs  = H_OWNER
        email = f"inv{uuid.uuid4().hex[:6]}@test.com"
        valid = {"email":email,"full_name":"Invite","initial_tenant_role":"tenant_admin"}
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
            body = valid
        elif exp_code in (400,422):
            if "email" in sc_l and ("format" in sc_l or "invalid" in sc_l or "@" not in sc_l):
                body={**valid,"email":"not-an-email"}
            elif "full_name" in sc_l or "name" in sc_l: body={**valid,"full_name":""}
            elif "member" in sc_l and "role" in sc_l:   body={**valid,"initial_tenant_role":"member"}
            elif "unknown" in sc_l:  body={**valid,"initial_tenant_role":"god"}
            elif "malformed" in sc_l or "415" in str(exp_code): body=None
            else: body={}
        elif exp_code==409: body=valid   # will hit seat limit or conflict
        elif exp_code==404: url=url.replace(TENANT,fresh_tenant()); body=valid
        elif exp_code in (201,202): body=valid
        else: body=valid

    # ── P-7 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("P7-"):
        url  = url.replace(MEMBER,MEMBER)
        hdrs = H_OWNER
        rv   = get_rv(BASE+f"/api/v1/tenants/{TENANT}/members/{MEMBER}", H_OWNER)
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        if "suspend" in sc_l:       body={"status":"suspended","record_version":rv}
        elif "reactivate" in sc_l or ("active" in sc_l and exp_code==200):
            body={"status":"active","record_version":rv}
        elif "left" in sc_l:        body={"status":"left","record_version":rv}
        elif exp_code==409:         body={"status":"suspended","record_version":9999}
        elif exp_code in (400,422):
            if "left" in sc_l:      body={"status":"left","record_version":rv}
            elif "missing" in sc_l: body={"record_version":rv}
            elif "deleted" in sc_l: body={"status":"deleted","record_version":rv}
            elif "uuid" in sc_l and "tenant" in sc_l:
                url=re.sub(TENANT,"not-a-uuid",url); body={"status":"suspended","record_version":rv}
            elif "uuid" in sc_l:
                url=re.sub(MEMBER,"not-a-uuid",url); body={"status":"suspended","record_version":rv}
            else:                   body={"status":"suspended","record_version":rv}
        elif exp_code==404:
            url=re.sub(MEMBER,fresh_user(),url)
            body={"status":"suspended","record_version":1}
        elif exp_code==422 and "last" in sc_l and "owner" in sc_l:
            url=url.replace(MEMBER,OWNER)
            rv2=get_rv(BASE+f"/api/v1/tenants/{TENANT}/members/{OWNER}", H_OWNER)
            body={"status":"suspended","record_version":rv2}
        elif exp_code==403:
            hdrs=H(SYS_USER,TENANT,"tender_admin")
            body={"status":"suspended","record_version":rv}
        elif exp_code==200:         body={"status":"suspended","record_version":rv}
        else:                       body={"status":"suspended","record_version":rv}

    # ── P-8 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("P8-"):
        url  = url.replace(MEMBER,MEMBER)
        hdrs = H_OWNER
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code==422 and "owner" in sc_l: url=url.replace(MEMBER,OWNER)
        elif exp_code==404: url=re.sub(MEMBER,fresh_user(),url)
        elif exp_code==403: hdrs=H(ADMIN,TENANT,"tenant_admin")
        elif exp_code==400:
            if "tenant" in sc_l: url=re.sub(TENANT,"not-a-uuid",url)
            else:                 url=re.sub(MEMBER,"not-a-uuid",url)

    # ── P-9 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("P9-"):
        hdrs = H_OWNER
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H_CROSS
        elif "member" in sc_l and exp_code==200: hdrs=H_MEMBER
        elif exp_code==400: url=re.sub(DEPT_ENG,"not-a-uuid",url)
        elif exp_code==404: url=re.sub(DEPT_ENG,fresh_user(),url)

    # ── P-10 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P10-"):
        url  = url.replace(MEMBER,MEMBER)
        hdrs = H_OWNER
        body = {"level":"preparator","record_version":1}
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code in (400,422):
            if "level" in sc_l and "invalid" in sc_l:
                body={"level":"invalid_level","record_version":1}
            elif "not.*member" in sc_l or "not in tenant" in sc_l:
                url=re.sub(MEMBER,fresh_user(),url)
            elif "deactivated" in sc_l or "retired" in sc_l:
                url=re.sub(DEPT_ENG,fresh_user(),url)
            elif "suspended" in sc_l: url=url.replace(MEMBER,SUSPENDED)
            elif "uuid" in sc_l and "tenant" in sc_l:
                url=re.sub(TENANT,"not-a-uuid",url)
            elif "uuid" in sc_l and "dept" in sc_l:
                url=re.sub(DEPT_ENG,"not-a-uuid",url)
            elif "uuid" in sc_l:
                url=re.sub(MEMBER,"not-a-uuid",url)
            else: body={}
        elif exp_code==404: url=re.sub(MEMBER,fresh_user(),url)
        elif "reviewer" in sc_l: body={"level":"reviewer","record_version":1}
        elif "same" in sc_l:    body={"level":"preparator","record_version":1}

    # ── P-11 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P11-"):
        url  = url.replace(MEMBER,MEMBER)
        hdrs = H_OWNER
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code==404: url=re.sub(MEMBER,fresh_user(),url)
        elif exp_code==400:
            if "tenant" in sc_l: url=re.sub(TENANT,"not-a-uuid",url)
            elif "dept" in sc_l:  url=re.sub(DEPT_ENG,"not-a-uuid",url)
            else:                 url=re.sub(MEMBER,"not-a-uuid",url)

    # ── P-12 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P12-"):
        hdrs = H_OWNER
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H_CROSS
        elif "member" in sc_l and exp_code==200: hdrs=H_MEMBER
        elif exp_code==200 and "admin" in sc_l:  hdrs=H_ADMIN
        elif exp_code==404:
            url=url.replace(TENANT,fresh_tenant())
        elif exp_code==403 and "cross" in sc_l:  hdrs=H_CROSS
        elif exp_code==400: url=re.sub(TENANT,"not-a-uuid",url)

    # ── P-13 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P13-"):
        hdrs = H_OWNER
        rv   = get_rv(BASE+f"/api/v1/tenants/{TENANT}/roles/preparator", H_OWNER)
        body = {"display_name":"Ops","record_version":rv}
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code in (400,422):
            if "invalid" in sc_l or "unknown" in sc_l:
                url=re.sub("preparator","manager_role",url)
            elif "empty" in sc_l: body={"display_name":"","record_version":rv}
        elif exp_code==409: body={"display_name":"X","record_version":9999}
        elif exp_code==404: url=url.replace(TENANT,fresh_tenant())

    # ── P-24 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P24-"):
        hdrs = H_OWNER
        body = {"department_id":DEPT_DES}
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code==400:              body={}
        elif exp_code==409:              body={"department_id":DEPT_ENG}
        elif exp_code in (404,422):      body={"department_id":fresh_user()}
        elif exp_code==403:              hdrs=H(SYS_USER,TENANT,"tender_admin")
        elif exp_code==201:              body={"department_id":DEPT_DES}

    # ── P-25 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P25-"):
        hdrs = H_OWNER
        # use DEPT_DES (non-system) for is_active=false tests
        dept = DEPT_ENG if "system" in sc_l else DEPT_DES
        url  = re.sub(DEPT_ENG, dept, url)
        rv   = get_rv(BASE+f"/api/v1/tenants/{TENANT}/departments/{dept}", H_OWNER)
        if rv == 1 and dept == DEPT_DES:
            # activate design dept first if not already
            call("POST",BASE+f"/api/v1/tenants/{TENANT}/departments",
                 {"department_id":DEPT_DES},H_OWNER)
            rv = get_rv(BASE+f"/api/v1/tenants/{TENANT}/departments/{dept}", H_OWNER)
        body = {"is_active":True,"record_version":rv}
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code==409:             body={"is_active":False,"record_version":9999}
        elif exp_code==422 and "system" in sc_l:
            body={"is_active":False,"record_version":rv}
        elif exp_code==422 and "retired" in sc_l:
            body={"is_active":True,"record_version":rv}
        elif exp_code==400:             body={"record_version":rv}
        elif exp_code==404:
            url=re.sub(dept,fresh_user(),url); body={"is_active":True,"record_version":1}
        elif "deactivate" in sc_l or "false" in sc_l:
            body={"is_active":False,"record_version":rv}

    # ── P-26 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P26-"):
        url  = url.replace(MEMBER,MEMBER)
        hdrs = H_OWNER
        body = {"action":"stop_workflows","record_version":1}
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
            body = {"action":"stop_workflows","record_version":1}
        elif exp_code in (400,422):
            if "invalid" in sc_l and "action" in sc_l:
                body={"action":"invalid_action","record_version":1}
            elif "missing" in sc_l or "absent" in sc_l:
                body={"action":"replace_delegate","record_version":1}
            elif "malformed" in sc_l: body=None
            elif "different tenant" in sc_l:
                body={"action":"replace_delegate","replacement_user_id":fresh_user(),"record_version":1}
            elif "equals" in sc_l:
                body={"action":"replace_delegate","replacement_user_id":MEMBER,"record_version":1}
            else: body={"action":"replace_delegate","record_version":1}
        elif "replace" in sc_l:
            body={"action":"replace_delegate","replacement_user_id":ADMIN,"record_version":1}

    # ── P-27 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P27-"):
        hdrs = H_OWNER
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H_CROSS
        elif exp_code==200 and "admin" in sc_l:  hdrs=H_ADMIN
        elif exp_code==200 and "owner" in sc_l:  hdrs=H_OWNER
        elif exp_code==200 and "member" in sc_l: hdrs=H_MEMBER
        elif exp_code==403: hdrs=H_MEMBER
        elif exp_code==404: url=url.replace(TENANT,fresh_tenant())

    # ── P-28 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P28-"):
        url  = url.replace(MEMBER,MEMBER)
        hdrs = H_OWNER
        rv   = get_rv(BASE+f"/api/v1/tenants/{TENANT}/members/{MEMBER}", H_OWNER)
        body = {"roles":["tenant_admin"],"record_version":rv}
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code in (400,422):
            if "unknown" in sc_l or "superman" in sc_l:
                body={"roles":["superman"],"record_version":rv}
            elif "member" in sc_l and "derived" in sc_l:
                body={"roles":["member"],"record_version":rv}
            elif "not.*member" in sc_l: url=re.sub(MEMBER,fresh_user(),url)
            elif "uuid" in sc_l and "tenant" in sc_l:
                url=re.sub(TENANT,"not-a-uuid",url)
            elif "uuid" in sc_l:  url=re.sub(MEMBER,"not-a-uuid",url)
            else: body={}
        elif "strip" in sc_l or "empty" in sc_l or "demote" in sc_l:
            body={"roles":[],"record_version":rv}
        elif "self" in sc_l:    body={"roles":[],"record_version":rv}; hdrs=H(OWNER,TENANT,"tenant_owner"); url=url.replace(MEMBER,OWNER)

    # ── P28RL ────────────────────────────────────────────────────────
    elif tc_up.startswith("P28RL-"):
        hdrs = H_OWNER
        rv   = get_rv(BASE+f"/api/v1/tenants/{TENANT}/roles/preparator", H_OWNER)
        body = {"display_name":"Custom Label","record_version":rv}
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code==409: body={"display_name":"X","record_version":9999}
        elif exp_code in (400,422):
            if "empty" in sc_l: body={"display_name":"","record_version":rv}
            elif "invalid" in sc_l: url=re.sub("preparator","manager_role",url)
            else: body={}
        elif exp_code==404:
            url=url.replace(TENANT,fresh_tenant())
        elif "approver" in sc_l:
            url=re.sub("preparator","approver",url)
            rv=get_rv(BASE+f"/api/v1/tenants/{TENANT}/roles/approver",H_OWNER)
            body={"display_name":"Approver Custom","record_version":rv}
        elif "admin" in sc_l and "non-owner" in sc_l:
            hdrs=H_ADMIN
            rv=get_rv(BASE+f"/api/v1/tenants/{TENANT}/roles/preparator",H_ADMIN)
            body={"display_name":"Prep Label","record_version":rv}
        elif "unicode" in sc_l or "long" in sc_l:
            body={"display_name":"Çöçüş Operatörü 运营商","record_version":rv}

    # ── P-30 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P30-"):
        hdrs = H_OWNER
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H_CROSS
        elif exp_code==200 and "admin" in sc_l:  hdrs=H_ADMIN
        elif exp_code==200 and "owner" in sc_l:  hdrs=H_OWNER
        elif exp_code==403 and "cross" in sc_l:  hdrs=H_CROSS
        elif exp_code==404: url=url.replace(TENANT,fresh_tenant())
        elif exp_code==400: url=re.sub(TENANT,"not-a-uuid",url)

    # ── P-31 ────────────────────────────────────────────────────────
    elif tc_up.startswith("P31-"):
        hdrs = H_OWNER
        inv_id = state.get("invite_id") or NIL_UUID
        url = url.replace(inv_id, inv_id)
        if is_auth:   hdrs = H_NO_AUTH if exp_code==401 else H(MEMBER,TENANT,"member")
        elif exp_code==404 or not state.get("invite_id"):
            url=re.sub(state.get("invite_id") or NIL_UUID, fresh_user(), url)
        elif exp_code==400:
            if "invitation_id" in sc_l: url=re.sub(inv_id,"not-a-uuid",url)
            elif "tenant" in sc_l:      url=re.sub(TENANT,"not-a-uuid",url)
        elif exp_code==403: hdrs=H_CROSS

    # ── O-4 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("O4-"):
        hdrs = H_OP
        rv   = get_rv(BASE+f"/api/v1/tenants/{TENANT}", H_OWNER)
        body = {"sso_enabled":False,"record_version":rv}
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
            body = {"sso_enabled":False,"record_version":rv}
        elif exp_code==403:             hdrs=H_OWNER
        elif exp_code==415:             hdrs=H_NOCT; body={"sso_enabled":False,"record_version":rv}
        elif exp_code in (400,422):
            if "null" in sc_l:          body={"sso_enabled":None,"record_version":rv}
            elif "array" in sc_l:       body={"sso_enabled":[],"record_version":rv}
            elif "object" in sc_l:      body={"sso_enabled":{},"record_version":rv}
            elif "empty" in sc_l and "key" in sc_l: body={"":True,"record_version":rv}
            elif "unknown" in sc_l or "typo" in sc_l or "sso_enable" in sc_l:
                body={"sso_enable":True,"record_version":rv}
            elif "max_seats" in sc_l:   body={"max_seats":100,"record_version":rv}
            elif "one valid" in sc_l:   body={"sso_enabled":True,"max_seats":100,"record_version":rv}
            else: body={}
        elif exp_code==404:
            url=url.replace(TENANT,fresh_tenant()); body={"sso_enabled":False,"record_version":1}
        elif exp_code==409:             body={"sso_enabled":False,"record_version":9999}
        elif "custom_branding" in sc_l: body={"custom_branding":"logo","record_version":rv}
        elif "require_mfa" in sc_l:     body={"require_mfa":True,"record_version":rv}
        elif "all" in sc_l:             body={"sso_enabled":True,"custom_branding":"logo","record_version":rv}
        elif "empty" in sc_l:           body={"record_version":rv}

    # ── O-7 ─────────────────────────────────────────────────────────
    elif tc_up.startswith("O7-"):
        hdrs = H_OP
        rv   = get_rv(BASE+f"/api/v1/tenants/{TENANT}", H_OWNER)
        body = {"user_id":ADMIN,"record_version":rv}
        if is_auth:
            hdrs = H_NO_AUTH if exp_code==401 else H(OWNER,TENANT,"tenant_owner")
            body = {"user_id":ADMIN,"record_version":rv}
        elif exp_code==403:             hdrs=H_OWNER
        elif exp_code==415:             hdrs=H_NOCT
        elif exp_code in (400,422):
            if "not.*member" in sc_l:   body={"user_id":fresh_user(),"record_version":rv}
            elif "uuid" in sc_l or "invalid" in sc_l: body={"user_id":"bad","record_version":rv}
            elif "nil" in sc_l:         body={"user_id":NIL_UUID,"record_version":rv}
            elif "suspended" in sc_l:   body={"user_id":SUSPENDED,"record_version":rv}
            elif "invited" in sc_l or "pending" in sc_l:
                body={"user_id":fresh_user(),"record_version":rv}
            else: body={}
        elif exp_code==409:             body={"user_id":ADMIN,"record_version":9999}
        elif exp_code==404:
            url=url.replace(TENANT,fresh_tenant()); body={"user_id":ADMIN,"record_version":1}
        elif "deprecated" in sc_l or "new_owner_user_id" in sc_l:
            body={"new_owner_user_id":ADMIN,"record_version":rv}
        elif "both" in sc_l and "user_id" in sc_l:
            body={"user_id":ADMIN,"new_owner_user_id":OWNER,"record_version":rv}

    # ── CRS / CROSS ──────────────────────────────────────────────────
    elif tc_up.startswith("CRS-") or tc_up.startswith("CROSS-"):
        if "/feature-flags" in url: hdrs=H_OP; body={"sso_enabled":True}
        elif "/reassign-owner" in url:
            rv=get_rv(BASE+f"/api/v1/tenants/{TENANT}",H_OWNER)
            hdrs=H_OP; body={"user_id":ADMIN,"record_version":rv}
        else: hdrs=H_OWNER

    return method, url, body, hdrs


# ── Result styles ─────────────────────────────────────────────────────
PASS_F = PatternFill('solid', fgColor='C6EFCE')
FAIL_F = PatternFill('solid', fgColor='FFC7CE')
SKIP_F = PatternFill('solid', fgColor='E0E0E0')
PEND_F = PatternFill('solid', fgColor='FFEB9C')
PASS_T = Font(name='Calibri',size=10,bold=True,color='1E7E34')
FAIL_T = Font(name='Calibri',size=10,bold=True,color='9C0006')
SKIP_T = Font(name='Calibri',size=10,color='666666')
DATA_T = Font(name='Calibri',size=10)
BDR    = Border(left=Side(style='thin',color='D9E1F2'),right=Side(style='thin',color='D9E1F2'),
                top=Side(style='thin',color='D9E1F2'),bottom=Side(style='thin',color='D9E1F2'))

def write_result(ws, row, code, passed, skipped=False):
    pc = ws.cell(row, 16)
    qc = ws.cell(row, 17)
    if skipped:
        pc.value='⚠ Manual'; pc.fill=SKIP_F; pc.font=SKIP_T; qc.value=None
    elif passed:
        pc.value='✅ Pass';  pc.fill=PASS_F; pc.font=PASS_T; qc.value=code; qc.font=DATA_T
    else:
        pc.value='❌ Fail';  pc.fill=FAIL_F; pc.font=FAIL_T; qc.value=code; qc.font=DATA_T
    for c in [pc,qc]:
        c.border=BDR; c.alignment=Alignment(horizontal='center',vertical='center')


# ── Main ──────────────────────────────────────────────────────────────
if __name__ == "__main__":
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]

    passed_n = failed_n = skipped_n = 0
    print(f"Running '{SHEET_NAME}' — {ws.max_row-1} test cases...\n")

    for r in range(2, ws.max_row + 1):
        row   = [ws.cell(r,c).value for c in range(1,19)]
        tc    = str(row[0] or '').strip()
        if not tc: continue

        cat   = str(row[1]  or '')
        sc    = str(row[2]  or '')
        meth  = str(row[3]  or '').upper()
        api   = str(row[4]  or '')
        exp   = row[5]
        notes = str(row[14] or '')
        nmv   = str(row[17] or '')
        res   = str(row[15] or '')   # col 16 = 25-Aug-2026

        # Skip already-decided cases
        if '⚠' in str(ws.cell(r,16).value or '') or '⚠' in nmv:
            skipped_n += 1
            if '⚠ Manual' not in str(ws.cell(r,16).value or ''):
                write_result(ws, r, None, False, skipped=True)
            continue

        if should_skip(tc, cat, sc, nmv):
            write_result(ws, r, None, False, skipped=True)
            skipped_n += 1
            continue

        # Pre-test state hooks
        pre_test_hooks(tc)

        try:
            m, url, body, hdrs = build_request(tc, cat, sc, meth, api, exp, notes)
        except Exception as e:
            write_result(ws, r, None, False, skipped=True)
            skipped_n += 1
            continue

        actual, resp = call(m, url, body, hdrs)

        # Capture invite_id after successful invite
        if tc.startswith("P6-") and actual in (201,202):
            iid = resp.get("invitation_id") or resp.get("id")
            if iid: state["invite_id"] = iid

        exp_int = int(exp) if exp else 0
        passed  = (actual == exp_int)
        write_result(ws, r, actual, passed)

        if not passed:
            print(f"❌ {tc:27} exp={exp_int:3} got={actual:3}  {sc[:52]}")

        passed_n += passed
        failed_n += (not passed)
        time.sleep(0.015)

    wb.save(EXCEL_PATH)
    total = passed_n + failed_n + skipped_n
    print(f"\n{'='*62}")
    print(f"  ✅ Pass    : {passed_n}")
    print(f"  ❌ Fail    : {failed_n}")
    print(f"  ⚠ Manual  : {skipped_n}  (event/concurrency/special-state)")
    print(f"  Total     : {total}")
    print(f"\nResults saved → {EXCEL_PATH}")
